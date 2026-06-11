import os
from datetime import datetime

from openpyxl import Workbook, load_workbook

from config import DATA_DIR

CROSS_CONNECT_HEADERS = [
    "Дата",
    "Сегмент",
    "Коммутатор",
    "Порт",
    "Номер СКС",
    "Комментарий",
]

PRINTER_HEADERS = [
    "Дата",
    "Домен",
    "Модель",
    "Серийный номер",
    "IP адрес",
    "Сберпечать ID",
    "Расположение",
]


def cross_connect_filename(floor: int) -> str:
    return os.path.join(DATA_DIR, f"Кроссовая_Этаж_{floor}.xlsx")


def printer_filename(domain: str) -> str:
    return os.path.join(DATA_DIR, f"Принтеры_{domain}.xlsx")


def append_printer_record(
    domain: str,
    model: str,
    serial: str,
    ip: str,
    sberprint_id: str,
    location: str,
) -> str:
    path = printer_filename(domain)

    if os.path.exists(path):
        workbook = load_workbook(path)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = f"Принтеры {domain}"
        worksheet.append(PRINTER_HEADERS)

    worksheet.append(
        [
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            domain,
            model,
            serial,
            ip,
            sberprint_id,
            location,
        ]
    )
    workbook.save(path)
    return path


def append_cross_connect_record(
    floor: int,
    segment: str,
    switch: int,
    port: int,
    scs_number: str,
    comment: str,
    ports_per_switch: int,
) -> str:
    path = cross_connect_filename(floor)

    if os.path.exists(path):
        workbook = load_workbook(path)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = f"Этаж {floor}"
        worksheet.append(CROSS_CONNECT_HEADERS)

    overall_port = (switch - 1) * ports_per_switch + port
    worksheet.append(
        [
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            segment,
            switch,
            overall_port,
            scs_number,
            comment or "",
        ]
    )
    workbook.save(path)
    return path


def list_data_files() -> list[str]:
    if not os.path.exists(DATA_DIR):
        return []
    return sorted(f for f in os.listdir(DATA_DIR) if f.endswith(".xlsx"))


def read_cross_connect_records(floor: int) -> list[tuple]:
    path = cross_connect_filename(floor)
    if not os.path.exists(path):
        return []

    workbook = load_workbook(path)
    worksheet = workbook.active

    records = []
    for row_idx in range(2, worksheet.max_row + 1):
        row = worksheet[row_idx]
        if row[0].value is None:
            continue
        _date, segment, switch, port, scs_number, comment = (cell.value for cell in row[:6])
        records.append((row_idx, segment, switch, port, scs_number, comment))
    return records


def get_cross_connect_record(floor: int, row_idx: int) -> tuple | None:
    for record in read_cross_connect_records(floor):
        if record[0] == row_idx:
            return record
    return None


def delete_cross_connect_record(floor: int, row_idx: int) -> bool:
    path = cross_connect_filename(floor)
    if not os.path.exists(path):
        return False

    workbook = load_workbook(path)
    worksheet = workbook.active
    if row_idx < 2 or row_idx > worksheet.max_row:
        return False

    worksheet.delete_rows(row_idx)
    workbook.save(path)
    return True


def update_cross_connect_record(floor: int, row_idx: int, scs_number: str, comment: str) -> bool:
    path = cross_connect_filename(floor)
    if not os.path.exists(path):
        return False

    workbook = load_workbook(path)
    worksheet = workbook.active
    if row_idx < 2 or row_idx > worksheet.max_row:
        return False

    worksheet.cell(row=row_idx, column=1, value=datetime.now().strftime("%Y-%m-%d %H:%M"))
    worksheet.cell(row=row_idx, column=5, value=scs_number)
    worksheet.cell(row=row_idx, column=6, value=comment or "")
    workbook.save(path)
    return True


def get_occupied_ports(floor: int, segment: str, switch: int, ports_per_switch: int) -> set[int]:
    occupied = set()
    for _row_idx, rec_segment, rec_switch, overall_port, _scs, _comment in read_cross_connect_records(floor):
        if rec_segment == segment and rec_switch == switch:
            port = overall_port - (switch - 1) * ports_per_switch
            if 1 <= port <= ports_per_switch:
                occupied.add(port)
    return occupied


def get_port_record(
    floor: int, segment: str, switch: int, port: int, ports_per_switch: int
) -> tuple | None:
    overall_port = (switch - 1) * ports_per_switch + port
    for record in read_cross_connect_records(floor):
        _row_idx, rec_segment, rec_switch, rec_overall_port, _scs, _comment = record
        if rec_segment == segment and rec_switch == switch and rec_overall_port == overall_port:
            return record
    return None
